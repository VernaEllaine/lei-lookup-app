"""API route handlers for the LEI Lookup web app."""

from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import sys
import uuid
from dataclasses import dataclass, field

import requests
from fastapi import APIRouter, File, Query, UploadFile
from fastapi.responses import StreamingResponse
from rapidfuzz import fuzz

from backend import database, gleif_local
from backend.cache import CacheManager
from backend.models import (
    CellUpdate,
    RowResult,
    Summary,
    UploadResponse,
    ValidateLeiResponse,
)
from backend.worker import run_lookup
from backend.validation_worker import run_validation

router = APIRouter(prefix="/api")


def _sniff_dialect(text: str) -> csv.Dialect:
    """Auto-detect CSV delimiter (comma, semicolon, tab, etc.)."""
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        return csv.excel  # default to comma


def _parse_upload(filename: str, content: bytes) -> tuple[list[str], list[dict]]:
    """Parse CSV or XLSX upload content. Returns (headers, rows)."""
    if filename.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None) or []
        headers = [str(c) if c is not None else '' for c in header_row]
        rows = [
            {headers[i]: (str(cell) if cell is not None else '') for i, cell in enumerate(row)}
            for row in rows_iter
        ]
        wb.close()
        return headers, rows
    else:
        text = content.decode('utf-8-sig')
        dialect = _sniff_dialect(text)
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = list(reader.fieldnames or [])
        rows = list(reader)
        return headers, rows

# Shared cache manager (global across sessions)
cache = CacheManager()

# Import detect helper
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from lei_lookup import _detect_name_column

LEI_FIELDS = [
    "lei", "lei_legal_name", "lei_jurisdiction", "lei_status",
    "lei_confidence", "lei_match_status", "lei_candidates",
]


# ---------------------------------------------------------------------------
# Per-session state
# ---------------------------------------------------------------------------

@dataclass
class SessionState:
    input_headers: list[str] = field(default_factory=list)
    rows: list[dict] = field(default_factory=list)
    name_column: str = ""
    running: bool = False


_sessions: dict[str, SessionState] = {}


def _get_session(session_id: str) -> SessionState | None:
    return _sessions.get(session_id)


def _build_summary(session_id: str) -> dict:
    s = database.compute_summary(session_id)
    s["cache_size"] = cache.size
    return s


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@router.post("/upload")
async def upload_csv(file: UploadFile = File(...)):
    session_id = uuid.uuid4().hex
    session = SessionState()
    _sessions[session_id] = session

    content = await file.read()
    session.input_headers, session.rows = _parse_upload(file.filename or '', content)

    detected = _detect_name_column(session.input_headers)
    session.name_column = detected or ""

    # Store rows in database
    database.clear_results(session_id)
    db_rows = []
    for row in session.rows:
        r = dict(row)
        r["_entity_name"] = r.get(session.name_column, "").strip() if session.name_column else ""
        db_rows.append(r)
    database.bulk_insert_results(db_rows, session_id)

    return {
        "session_id": session_id,
        "headers": session.input_headers,
        "row_count": len(session.rows),
        "detected_column": detected,
    }


@router.get("/lookup")
async def start_lookup(
    column: str = Query(...),
    session_id: str = Query(...),
):
    session = _get_session(session_id)
    if session is None:
        return StreamingResponse(
            iter(['data: {"error": "Invalid session"}\n\n']),
            media_type="text/event-stream",
        )

    if session.running:
        return StreamingResponse(
            iter(['data: {"error": "Lookup already running"}\n\n']),
            media_type="text/event-stream",
        )

    if not session.rows:
        return StreamingResponse(
            iter(['data: {"error": "No CSV uploaded"}\n\n']),
            media_type="text/event-stream",
        )

    session.name_column = column
    session.running = True

    queue: asyncio.Queue = asyncio.Queue()

    async def event_generator():
        total = len(session.rows)
        received = 0
        task = asyncio.create_task(
            run_lookup(session.rows, column, cache, queue, session_id)
        )

        try:
            while received < total:
                try:
                    progress = await asyncio.wait_for(queue.get(), timeout=0.5)
                except asyncio.TimeoutError:
                    if task.done() and queue.empty():
                        break
                    if not task.done():
                        yield ":\n\n"  # keepalive
                    continue

                received += 1
                progress.done = received == total
                data = progress.model_dump_json()
                yield f"data: {data}\n\n"
        except asyncio.CancelledError:
            task.cancel()
            raise
        finally:
            session.running = False
            summary = _build_summary(session_id)
            yield f"event: summary\ndata: {json.dumps(summary)}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/results")
async def get_results(
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=1, le=200000),
    status: str | None = Query(None),
    session_id: str = Query(""),
):
    result = database.get_results_page(page, page_size, status, session_id=session_id)
    return {
        "rows": result["rows"],
        "total": result["total"],
        "page": result["page"],
        "page_size": result["page_size"],
        "summary": _build_summary(session_id),
    }


@router.put("/results/{row_id}")
async def update_cell(row_id: int, update: CellUpdate):
    row_data = database.get_result_by_id(row_id)
    if row_data is None:
        return {"error": "Invalid index"}

    field_map = {
        "lei": "lei",
        "legal_name": "legal_name",
        "jurisdiction": "jurisdiction",
        "status": "status",
        "confidence": "confidence",
    }

    if update.field not in field_map:
        return {"error": f"Invalid field: {update.field}"}

    attr = field_map[update.field]
    old_value = row_data.get(attr, "")

    db_update = {attr: update.value}

    # If LEI changed, don't mark reviewed yet — frontend should call validate-lei
    if update.field != "lei" or update.value == old_value:
        db_update["match_status"] = "REVIEWED"
        # Cache the reviewed row
        entity_name = row_data.get("entity_name", "")
        if entity_name:
            cache.cache_row(entity_name, {
                "lei": update.value if attr == "lei" else row_data.get("lei", ""),
                "legal_name": update.value if attr == "legal_name" else row_data.get("legal_name", ""),
                "jurisdiction": update.value if attr == "jurisdiction" else row_data.get("jurisdiction", ""),
                "status": update.value if attr == "status" else row_data.get("status", ""),
                "confidence": update.value if attr == "confidence" else row_data.get("confidence", ""),
                "match_status": "REVIEWED",
            })
            cache.flush()

    database.update_result_by_id(row_id, db_update)
    updated_row = database.get_result_by_id(row_id)
    # Derive session_id from the row's session
    sid = _session_id_for_row(row_id)
    return {"row": updated_row, "summary": _build_summary(sid)}


@router.put("/results/{row_id}/validate-lei")
async def validate_lei(row_id: int):
    row_data = database.get_result_by_id(row_id)
    if row_data is None:
        return ValidateLeiResponse(valid=False, message="Invalid index")

    lei_code = row_data.get("lei", "")
    entity_name = row_data.get("entity_name", "")

    if not lei_code:
        return ValidateLeiResponse(valid=False, message="No LEI to validate")

    try:
        local = gleif_local.get_record(lei_code)
        if local is not None:
            legal_name = local["legal_name"]
            jurisdiction = local["jurisdiction"]
            status = local["status"]
            reg_status = local["registration_status"]
        else:
            resp = requests.get(
                f"https://api.gleif.org/api/v1/lei-records/{lei_code}",
                timeout=15,
            )
            if resp.status_code != 200:
                return ValidateLeiResponse(
                    valid=False,
                    message=f"LEI '{lei_code}' not found (HTTP {resp.status_code}).",
                )

            data = resp.json().get("data", {})
            attr = data.get("attributes", {}).get("entity", {})
            legal_name = attr.get("legalName", {}).get("name", "")
            jurisdiction = attr.get("jurisdiction", "")
            status = attr.get("status", "")
            reg = data.get("attributes", {}).get("registration", {})
            reg_status = reg.get("status", "")

        if status != "ACTIVE" or reg_status != "ISSUED":
            return ValidateLeiResponse(
                valid=False,
                message=f"LEI '{lei_code}' is not active (entity: {status}, registration: {reg_status}).",
            )

        score = (
            fuzz.token_sort_ratio(entity_name.lower(), legal_name.lower())
            if entity_name and legal_name
            else 0
        )
        needs_confirmation = score < 50

        database.update_result_by_id(row_id, {
            "legal_name": legal_name,
            "jurisdiction": jurisdiction,
            "status": status,
            "match_status": "REVIEWED",
        })

        if entity_name:
            cache.cache_row(entity_name, {
                "lei": lei_code,
                "legal_name": legal_name,
                "jurisdiction": jurisdiction,
                "status": status,
                "confidence": row_data.get("confidence", ""),
                "match_status": "REVIEWED",
            })
            cache.flush()

        return ValidateLeiResponse(
            valid=True,
            legal_name=legal_name,
            jurisdiction=jurisdiction,
            status=status,
            needs_confirmation=needs_confirmation,
            message="LEI validated successfully." if not needs_confirmation
            else f"GLEIF legal name '{legal_name}' does not closely match entity '{entity_name}'.",
        )
    except Exception as exc:
        return ValidateLeiResponse(valid=False, message=f"Validation error: {exc}")


@router.post("/confirm-all")
async def confirm_all(session_id: str = Query("")):
    count = database.confirm_reviewed_rows(session_id)
    return {"confirmed": count, "summary": _build_summary(session_id)}


@router.get("/export")
async def export_csv(session_id: str = Query("")):
    session = _get_session(session_id)
    all_rows = database.get_all_results_for_export(session_id)
    if not all_rows:
        return StreamingResponse(
            iter([""]),
            media_type="text/csv",
        )

    input_headers = session.input_headers if session else []
    output = io.StringIO()
    output_headers = list(input_headers) + LEI_FIELDS
    writer = csv.DictWriter(output, fieldnames=output_headers, extrasaction="ignore")
    writer.writeheader()

    for row_data in all_rows:
        original = row_data.get("original", {})
        out = dict(original)
        out["lei"] = row_data["lei"]
        out["lei_legal_name"] = row_data["legal_name"]
        out["lei_jurisdiction"] = row_data["jurisdiction"]
        out["lei_status"] = row_data["status"]
        out["lei_confidence"] = row_data["confidence"]
        out["lei_match_status"] = row_data["match_status"]
        out["lei_candidates"] = row_data["candidates"]
        writer.writerow(out)

    csv_content = output.getvalue()
    return StreamingResponse(
        iter([csv_content]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=lei_results.csv"},
    )


@router.get("/export-xlsx")
async def export_xlsx(session_id: str = Query("")):
    session = _get_session(session_id)
    all_rows = database.get_all_results_for_export(session_id)
    if not all_rows:
        return StreamingResponse(
            iter([b""]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    import openpyxl

    input_headers = session.input_headers if session else []
    output_headers = list(input_headers) + LEI_FIELDS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(output_headers)

    for row_data in all_rows:
        original = row_data.get("original", {})
        out = dict(original)
        out["lei"] = row_data["lei"]
        out["lei_legal_name"] = row_data["legal_name"]
        out["lei_jurisdiction"] = row_data["jurisdiction"]
        out["lei_status"] = row_data["status"]
        out["lei_confidence"] = row_data["confidence"]
        out["lei_match_status"] = row_data["match_status"]
        out["lei_candidates"] = row_data["candidates"]
        ws.append([out.get(h, "") for h in output_headers])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=lei_results.xlsx"},
    )


@router.get("/cache/summary")
async def cache_summary(session_id: str = Query("")):
    return {"cache_size": cache.size, "summary": _build_summary(session_id)}


@router.delete("/cache")
async def clear_cache(session_id: str = Query("")):
    cache.clear()
    return {"message": "Cache cleared.", "summary": _build_summary(session_id)}


# ---------------------------------------------------------------------------
# Validation session state
# ---------------------------------------------------------------------------

@dataclass
class ValidationSessionState:
    input_headers: list[str] = field(default_factory=list)
    rows: list[dict] = field(default_factory=list)
    entity_column: str = ""
    lei_column: str = ""
    running: bool = False


_validation_sessions: dict[str, ValidationSessionState] = {}


def _detect_lei_column(headers: list[str]) -> str | None:
    """Guess which column holds the LEI code."""
    candidates = ["lei", "lei_code", "lei_number", "legal_entity_identifier"]
    lower_headers = {h.lower().strip(): h for h in headers}
    for c in candidates:
        if c in lower_headers:
            return lower_headers[c]
    # Fallback: any header containing "lei"
    for h in headers:
        if "lei" in h.lower():
            return h
    return None


# ---------------------------------------------------------------------------
# Validation Routes
# ---------------------------------------------------------------------------

@router.post("/validate/upload")
async def validate_upload_csv(file: UploadFile = File(...)):
    session_id = uuid.uuid4().hex
    session = ValidationSessionState()
    _validation_sessions[session_id] = session

    content = await file.read()
    session.input_headers, session.rows = _parse_upload(file.filename or '', content)

    detected_entity = _detect_name_column(session.input_headers)
    detected_lei = _detect_lei_column(session.input_headers)
    session.entity_column = detected_entity or ""
    session.lei_column = detected_lei or ""

    # Store rows in database
    database.clear_validation_results(session_id)
    db_rows = []
    for row in session.rows:
        r = dict(row)
        r["_entity_name"] = r.get(session.entity_column, "").strip() if session.entity_column else ""
        r["_provided_lei"] = r.get(session.lei_column, "").strip() if session.lei_column else ""
        db_rows.append(r)
    database.bulk_insert_validation_results(db_rows, session_id)

    return {
        "session_id": session_id,
        "headers": session.input_headers,
        "row_count": len(session.rows),
        "detected_entity_column": detected_entity,
        "detected_lei_column": detected_lei,
    }


@router.get("/validate/run")
async def start_validation(
    entity_column: str = Query(...),
    lei_column: str = Query(...),
    session_id: str = Query(...),
):
    session = _validation_sessions.get(session_id)
    if session is None:
        return StreamingResponse(
            iter(['data: {"error": "Invalid session"}\n\n']),
            media_type="text/event-stream",
        )

    if session.running:
        return StreamingResponse(
            iter(['data: {"error": "Validation already running"}\n\n']),
            media_type="text/event-stream",
        )

    if not session.rows:
        return StreamingResponse(
            iter(['data: {"error": "No CSV uploaded"}\n\n']),
            media_type="text/event-stream",
        )

    session.entity_column = entity_column
    session.lei_column = lei_column
    session.running = True

    # Re-store rows with correct columns in DB
    database.clear_validation_results(session_id)
    db_rows = []
    for row in session.rows:
        r = dict(row)
        r["_entity_name"] = r.get(entity_column, "").strip()
        r["_provided_lei"] = r.get(lei_column, "").strip()
        db_rows.append(r)
    database.bulk_insert_validation_results(db_rows, session_id)

    queue: asyncio.Queue = asyncio.Queue()

    async def event_generator():
        total = len(session.rows)
        received = 0
        task = asyncio.create_task(
            run_validation(session.rows, entity_column, lei_column, queue, session_id)
        )

        try:
            while received < total:
                try:
                    progress = await asyncio.wait_for(queue.get(), timeout=0.5)
                except asyncio.TimeoutError:
                    if task.done() and queue.empty():
                        break
                    if not task.done():
                        yield ":\n\n"
                    continue

                received += 1
                progress.done = received == total
                data = progress.model_dump_json()
                yield f"data: {data}\n\n"
        except asyncio.CancelledError:
            task.cancel()
            raise
        finally:
            session.running = False
            summary = database.compute_validation_summary(session_id)
            yield f"event: summary\ndata: {json.dumps(summary)}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/validate/results")
async def get_validation_results(
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=1, le=200000),
    flag: str | None = Query(None),
    session_id: str = Query(""),
):
    result = database.get_validation_results_page(page, page_size, flag, session_id=session_id)
    return {
        "rows": result["rows"],
        "total": result["total"],
        "page": result["page"],
        "page_size": result["page_size"],
        "summary": database.compute_validation_summary(session_id),
    }


@router.put("/validate/accept")
async def accept_suggested_lei(
    session_id: str = Query(...),
    row_index: int = Query(...),
):
    """Accept a suggested LEI: copy suggestion → provided, mark OK."""
    conn = database._get_conn()
    row = conn.execute(
        "SELECT suggested_lei FROM validation_results WHERE session_id = ? AND row_index = ?",
        (session_id, row_index),
    ).fetchone()
    if row is None:
        return {"error": "Row not found"}

    suggested = row["suggested_lei"] or ""
    if not suggested:
        return {"error": "No suggested LEI to accept"}

    database.update_validation_result(row_index, {
        "provided_lei": suggested,
        "entity_status": "ACTIVE",
        "registration_status": "ISSUED",
        "flag": "OK",
        "suggested_lei": "",
        "suggested_legal_name": "",
        "suggested_confidence": "",
    }, session_id)

    return {"summary": database.compute_validation_summary(session_id)}


@router.get("/validate/export")
async def export_validation_csv(session_id: str = Query("")):
    all_rows = database.get_all_validation_results_for_export(session_id)
    if not all_rows:
        return StreamingResponse(
            iter([""]),
            media_type="text/csv",
        )

    output = io.StringIO()
    output_headers = [
        "entity_name", "provided_lei", "entity_status", "registration_status",
        "flag", "suggested_lei", "suggested_legal_name", "suggested_confidence",
    ]
    writer = csv.DictWriter(output, fieldnames=output_headers, extrasaction="ignore")
    writer.writeheader()

    for row_data in all_rows:
        writer.writerow({k: row_data.get(k, "") for k in output_headers})

    csv_content = output.getvalue()
    return StreamingResponse(
        iter([csv_content]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=lei_validation_results.csv"},
    )


@router.get("/manual-lookup")
async def manual_lookup(names: str = Query(...)):
    """Look up company names directly without a CSV upload.

    Accepts a comma- or newline-separated list of company names.  Each name is
    run through lookup_lei() and classified, returning the same fields that the
    CSV batch processor writes to each row.
    """
    from lei_lookup import lookup_lei as _lookup, _classify, _best_match, GleifAPIError

    name_list = [n.strip() for n in names.replace("\n", ",").split(",") if n.strip()]
    if not name_list:
        return {"results": []}

    results: list[dict] = []
    for name in name_list:
        try:
            result = _lookup(name)
            status = _classify(result)
            best = _best_match(result)

            if status == "AUTO-MATCHED" and best:
                results.append({
                    "query": name,
                    "lei": best["lei"],
                    "legal_name": best["legal_name"],
                    "jurisdiction": best["jurisdiction"],
                    "status": best["status"],
                    "confidence": best["confidence"],
                    "match_status": "AUTO-MATCHED",
                    "candidates": "",
                    "error": "",
                })
            elif status == "REVIEW NEEDED":
                top = result["results"][0]
                candidates = [
                    f"{r['legal_name']} | {r['lei']} | {r['jurisdiction']} | {r['confidence']}"
                    for r in result["results"]
                ]
                results.append({
                    "query": name,
                    "lei": top["lei"],
                    "legal_name": top["legal_name"],
                    "jurisdiction": top["jurisdiction"],
                    "status": top["status"],
                    "confidence": top["confidence"],
                    "match_status": "REVIEW NEEDED",
                    "candidates": "; ".join(candidates),
                    "error": "",
                })
            else:
                results.append({
                    "query": name,
                    "lei": "", "legal_name": "", "jurisdiction": "",
                    "status": "", "confidence": "",
                    "match_status": "NO MATCH",
                    "candidates": "", "error": "",
                })
        except Exception as exc:
            results.append({
                "query": name,
                "lei": "", "legal_name": "", "jurisdiction": "",
                "status": "", "confidence": "",
                "match_status": "ERROR",
                "candidates": "", "error": str(exc),
            })

    return {"results": results}


@router.post("/validate/manual")
async def manual_validate(entries: list[dict]):
    """Validate LEI codes directly without a CSV upload.

    Accepts a JSON array of {"entity_name": "...", "lei": "..."} objects.
    Each LEI is checked against the GLEIF API and flagged OK / LAPSED /
    INVALID / NOT_FOUND / ERROR.
    """
    from lei_lookup import validate_single_lei as _validate

    results: list[dict] = []
    for entry in entries:
        entity_name = entry.get("entity_name", "").strip()
        lei = entry.get("lei", "").strip().upper()

        if not lei:
            results.append({
                "entity_name": entity_name, "lei": lei,
                "entity_status": "", "registration_status": "",
                "legal_name": "", "jurisdiction": "",
                "flag": "INVALID", "error": "No LEI provided",
            })
            continue

        v = _validate(lei)
        results.append({
            "entity_name": entity_name,
            "lei": lei,
            "entity_status": v["entity_status"],
            "registration_status": v["registration_status"],
            "legal_name": v["legal_name"],
            "jurisdiction": v["jurisdiction"],
            "flag": v["flag"],
            "error": "",
        })

    return {"results": results}


@router.get("/isin-lookup")
async def isin_lookup(isins: str = Query(...)):
    """Look up ISIN codes via the GLEIF lei-records filter and return LEI + entity details."""
    isin_list = [part.strip().upper() for part in isins.replace("\n", ",").split(",") if part.strip()]
    if not isin_list:
        return {"results": []}

    results: list[dict] = []
    seen_leis: set[str] = set()

    for isin in isin_list:
        local = [
            rec for rec in map(gleif_local.get_record, gleif_local.leis_for_isin(isin))
            if rec is not None
        ]
        if local:
            for rec in local:
                if rec["lei"] in seen_leis:
                    continue
                seen_leis.add(rec["lei"])
                results.append({
                    "isin": isin,
                    "lei": rec["lei"],
                    "legal_name": rec["legal_name"],
                    "country": rec["jurisdiction"],
                    "entity_status": rec["status"],
                    "registration_status": rec["registration_status"],
                    "error": "",
                })
            continue

        try:
            resp = requests.get(
                "https://api.gleif.org/api/v1/lei-records",
                params={"filter[isin]": isin, "page[size]": "10"},
                headers={"Accept": "application/vnd.api+json"},
                timeout=15,
            )
            if resp.status_code != 200:
                results.append({
                    "isin": isin, "lei": "", "legal_name": "", "country": "",
                    "entity_status": "", "registration_status": "",
                    "error": f"ISIN lookup failed (HTTP {resp.status_code})",
                })
                continue

            records = resp.json().get("data", [])
            if not records:
                results.append({
                    "isin": isin, "lei": "", "legal_name": "", "country": "",
                    "entity_status": "", "registration_status": "",
                    "error": "No LEI found for this ISIN",
                })
                continue

            for rec in records:
                lei = rec.get("id", "")
                if not lei or lei in seen_leis:
                    continue
                seen_leis.add(lei)

                attrs = rec.get("attributes", {})
                entity = attrs.get("entity", {})
                registration = attrs.get("registration", {})

                results.append({
                    "isin": isin,
                    "lei": lei,
                    "legal_name": entity.get("legalName", {}).get("name", ""),
                    "country": entity.get("jurisdiction", ""),
                    "entity_status": entity.get("status", ""),
                    "registration_status": registration.get("status", ""),
                    "error": "",
                })

        except Exception as exc:
            results.append({
                "isin": isin, "lei": "", "legal_name": "", "country": "",
                "entity_status": "", "registration_status": "",
                "error": str(exc),
            })

    return {"results": results}


@router.get("/gleif-local/status")
async def gleif_local_status():
    """Report whether a local GLEIF copy is loaded and how fresh it is."""
    return gleif_local.status()


@router.post("/isin-lookup/export-xlsx")
async def export_isin_xlsx(payload: dict):
    """Accept ISIN lookup results as JSON and return an XLSX file."""
    import openpyxl

    results = payload.get("results", [])
    headers = [
        "ISIN", "LEI", "Legal Entity Name", "Country",
        "Entity Status", "Registration Status", "Error",
    ]
    keys = [
        "isin", "lei", "legal_name", "country",
        "entity_status", "registration_status", "error",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in results:
        ws.append([row.get(k, "") for k in keys])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=isin_lookup_results.xlsx"},
    )


@router.get("/validate/export-xlsx")
async def export_validation_xlsx(session_id: str = Query("")):
    all_rows = database.get_all_validation_results_for_export(session_id)
    if not all_rows:
        return StreamingResponse(
            iter([b""]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    import openpyxl

    output_headers = [
        "entity_name", "provided_lei", "entity_status", "registration_status",
        "flag", "suggested_lei", "suggested_legal_name", "suggested_confidence",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(output_headers)

    for row_data in all_rows:
        ws.append([row_data.get(h, "") for h in output_headers])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=lei_validation_results.xlsx"},
    )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _session_id_for_row(row_id: int) -> str:
    """Look up the session_id for a given row primary key."""
    conn = database._get_conn()
    row = conn.execute("SELECT session_id FROM results WHERE id = ?", (row_id,)).fetchone()
    return row["session_id"] if row else ""
